# -*- coding: utf-8 -*-
"""将三份 Markdown 规格说明书转换为 Word(.docx) 与 PDF，中文友好。"""
import re, os

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, Preformatted, HRFlowable)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

BASE = r"D:\code\otherProjects\22_StudyWork"
FILES = [
    "阳光学习能量站-需求规格说明书",
    "阳光学习能量站-概要设计说明书",
    "阳光学习能量站-详细设计说明书",
]
CJK = "STSong-Light"
pdfmetrics.registerFont(UnicodeCIDFont(CJK))

# ---------------- Markdown 解析 ----------------
def parse_blocks(lines):
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i].rstrip("\n")
        s = line.strip()
        if s == "":
            i += 1; continue
        if s.startswith("```"):
            i += 1; code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip("\n")); i += 1
            i += 1
            blocks.append({"type": "code", "text": "\n".join(code)}); continue
        if re.match(r"^-{3,}$", s) or re.match(r"^\*{3,}$", s):
            blocks.append({"type": "hr"}); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2).strip()}); i += 1; continue
        if line.lstrip().startswith(">"):
            q = []
            while i < n and lines[i].lstrip().startswith(">"):
                q.append(re.sub(r"^>\s?", "", lines[i].rstrip("\n"))); i += 1
            blocks.append({"type": "quote", "text": "\n".join(q)}); continue
        if s.startswith("|"):
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip()); i += 1
            def split_row(r):
                r = r.strip()
                if r.startswith("|"): r = r[1:]
                if r.endswith("|"): r = r[:-1]
                parts = re.split(r'(?<!\\)\|', r)
                return [p.strip().replace('\\|', '|') for p in parts]
            if len(tbl) >= 2:
                headers = split_row(tbl[0])
                rows = [split_row(r) for r in tbl[2:]]
                blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue
        if re.match(r"^[-*]\s+", line):
            items = []
            while i < n and re.match(r"^[-*]\s+", lines[i].rstrip("\n")):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].rstrip("\n"))); i += 1
            blocks.append({"type": "ul", "items": items}); continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while i < n and re.match(r"^\d+\.\s+", lines[i].rstrip("\n")):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].rstrip("\n"))); i += 1
            blocks.append({"type": "ol", "items": items}); continue
        para = []
        while i < n:
            ls = lines[i].rstrip("\n")
            if (ls.strip() == "" or ls.strip().startswith("```") or re.match(r"^#{1,6}\s", ls)
                    or ls.strip().startswith("|") or ls.lstrip().startswith(">")
                    or re.match(r"^[-*]\s+", ls) or re.match(r"^\d+\.\s+", ls)
                    or re.match(r"^-{3,}$", ls.strip())):
                break
            para.append(ls); i += 1
        blocks.append({"type": "p", "text": " ".join(x.strip() for x in para)})
    return blocks

def inline_segments(text):
    out, cur, b, c, i = [], "", False, False, 0
    while i < len(text):
        if text[i:i+2] == "**":
            if cur: out.append((cur, b, c)); cur = ""
            b = not b; i += 2; continue
        if text[i] == "`":
            if cur: out.append((cur, b, c)); cur = ""
            c = not c; i += 1; continue
        cur += text[i]; i += 1
    if cur: out.append((cur, b, c))
    return out

# ---------------- DOCX ----------------
def set_cjk(run):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "SimSun")

def add_runs(para, text):
    for txt, b, c in inline_segments(text):
        r = para.add_run(txt)
        if b: r.bold = True
        if c: r.font.name = "Courier New"
        set_cjk(r)

def gen_docx(path, blocks):
    doc = Document()
    for b in blocks:
        t = b["type"]
        if t == "h":
            p = doc.add_heading(level=min(b["level"], 9))
            add_runs(p, b["text"])
        elif t == "p":
            p = doc.add_paragraph(); add_runs(p, b["text"])
        elif t == "quote":
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, b["text"])
        elif t == "ul":
            for it in b["items"]:
                p = doc.add_paragraph(style="List Bullet"); add_runs(p, it)
        elif t == "ol":
            for it in b["items"]:
                p = doc.add_paragraph(style="List Number"); add_runs(p, it)
        elif t == "code":
            p = doc.add_paragraph(); r = p.add_run(b["text"])
            r.font.name = "Courier New"; r.font.size = Pt(9); set_cjk(r)
        elif t == "hr":
            doc.add_paragraph("─" * 28)
        elif t == "table":
            tbl = doc.add_table(rows=1, cols=len(b["headers"]))
            tbl.style = "Light Grid Accent 1"
            for j, h in enumerate(b["headers"]):
                cell = tbl.rows[0].cells[j]; cell.text = ""
                r = cell.paragraphs[0].add_run(strip_md(h)); r.bold = True; set_cjk(r)
            for row in b["rows"]:
                row = (list(row) + [""] * len(b["headers"]))[:len(b["headers"])]
                cells = tbl.add_row().cells
                for j, c in enumerate(row):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], c)
    doc.save(path)

def strip_md(text):
    return re.sub(r"\*\*(.*?)\*\*", r"\1", text).replace("`", "")

# ---------------- PDF ----------------
def rl_inline(text):
    out = []
    for txt, b, c in inline_segments(text):
        e = txt.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if c: e = '<font name="Courier">%s</font>' % e
        if b: e = "<b>%s</b>" % e
        out.append(e)
    return "".join(out)

def gen_pdf(path, blocks):
    base = ParagraphStyle("base", fontName=CJK, fontSize=10.5, leading=16)
    para = ParagraphStyle("para", parent=base, firstLineIndent=18, spaceAfter=4)
    h1 = ParagraphStyle("h1", parent=base, fontSize=17, leading=22, spaceBefore=10, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontSize=14, leading=19, spaceBefore=8, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=base, fontSize=12, leading=16, spaceBefore=6, spaceAfter=3)
    quote = ParagraphStyle("quote", parent=base, leftIndent=10, textColor=colors.HexColor("#6b8a78"))
    code = ParagraphStyle("code", parent=base, fontName="Courier", fontSize=8.5,
                          leading=11, backColor=colors.whitesmoke, borderPadding=4)
    th = ParagraphStyle("th", parent=base, fontSize=9, leading=12, textColor=colors.white)
    td = ParagraphStyle("td", parent=base, fontSize=8.5, leading=11)
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=18*mm, bottomMargin=16*mm,
                            leftMargin=18*mm, rightMargin=18*mm)
    flow = []
    for b in blocks:
        t = b["type"]
        if t == "h":
            st = {1: h1, 2: h2, 3: h3}.get(b["level"], para)
            flow.append(Paragraph(rl_inline(b["text"]), st))
        elif t == "p":
            flow.append(Paragraph(rl_inline(b["text"]), para))
        elif t == "quote":
            flow.append(Paragraph(rl_inline(b["text"]), quote))
        elif t == "ul":
            for it in b["items"]:
                flow.append(Paragraph("• " + rl_inline(it), base))
        elif t == "ol":
            for i, it in enumerate(b["items"], 1):
                flow.append(Paragraph("%d. %s" % (i, rl_inline(it)), base))
        elif t == "code":
            flow.append(Preformatted(b["text"], code))
        elif t == "hr":
            flow.append(HRFlowable(width="100%", thickness=0.6, color=colors.grey,
                                  spaceBefore=4, spaceAfter=4))
        elif t == "table":
            data = [[Paragraph(rl_inline(h), th) for h in b["headers"]]]
            for row in b["rows"]:
                row = (list(row) + [""] * len(b["headers"]))[:len(b["headers"])]
                data.append([Paragraph(rl_inline(c), td) for c in row])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e8b57")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1fbf4")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            flow.append(tbl)
    doc.build(flow)

# ---------------- 主流程 ----------------
if __name__ == "__main__":
    for name in FILES:
        src = os.path.join(BASE, name + ".md")
        with open(src, "r", encoding="utf-8") as f:
            blocks = parse_blocks(f.readlines())
        dx = os.path.join(BASE, name + ".docx")
        pf = os.path.join(BASE, name + ".pdf")
        gen_docx(dx, blocks)
        gen_pdf(pf, blocks)
        print("OK:", name, "->", os.path.exists(dx) and os.path.exists(pf))
